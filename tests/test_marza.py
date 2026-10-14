"""Moduł MARŻA – podłączenie do silnika bez zmiany obliczeń (wysyłka przez serwer, chmura, przywrócenie smtplib,
tabela miesięcy dla zakładki WYNIKI)."""
import contextlib
import datetime
import io
import os
import re
import sys
import tempfile
import types
import zipfile

from _wspolne import temp_data  # noqa: F401 – ustawia ścieżkę do app/

MIESIACE = "/rest/v1/marza_months?on_conflict=branch,year,month"


class _Sesja:
    email = "test@example.com"

    def __init__(self, fail_paths=(), fail_upload=False):
        self.calls = []
        self.fail_paths, self.fail_upload = fail_paths, fail_upload

    def call(self, m, p, b=None, prefer=None):
        self.calls.append((m, p, b, prefer))
        if any(p.startswith(f) for f in self.fail_paths):
            import cloud
            raise cloud.CloudError("Brak połączenia z serwerem.")
        return {"ok": True}

    def upload(self, bucket, path, data, ctype):
        self.calls.append(("UP", path, len(data)))
        if self.fail_upload:
            import cloud
            raise cloud.CloudError("Nie zapisano pliku w chmurze.")


def _silnik(folder, fail=False):
    eng = types.SimpleNamespace(XLSX_NAME="k.xlsx", EMAIL_NADAWCA="", EMAIL_ODBIORCY=[], LOG_NAME="log.txt")
    eng.extract_text = lambda p: "tekst"
    eng.parse_report = lambda t, n: {"label": "Październik 2026", "margin_netto": 12345.5, "start_date": "2026-10-01",
                                     "end_date": "2026-10-13", "source": n}
    eng._refresh_podglad = lambda x, log: ("Październik 2026", 300000, "120000.5", set())
    eng._wczytaj_haslo_email = lambda: None

    def main():
        pdf = os.path.join(folder, "r.pdf")
        eng.parse_report(eng.extract_text(pdf), "r.pdf")
        eng._refresh_podglad("k", [])
        import smtplib   # tak jak oryginalny silnik – import w środku funkcji
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=20) as s:
            s.starttls()
            s.login(eng.EMAIL_NADAWCA, eng._wczytaj_haslo_email())
            for o in eng.EMAIL_ODBIORCY:
                s.sendmail(eng.EMAIL_NADAWCA, [o], "MIME...")
        if fail:
            raise RuntimeError("awaria silnika")
    eng.main = main
    return eng


def _kalkulator(path):
    """Kalkulator jak u właściciela (arkusz „Kalkulator marży” z tabelą od nagłówka „Miesiąc” + „Święta PL”).
    Listopad: cel to formuła z wynikiem zapisanym przez Excel; Grudzień: formuła bez wyniku (po zapisie przez
    openpyxl); „Razem”: wiersz bez miesiąca; Październik 2026 dwa razy (silnik pisze do ostatniego)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kalkulator marży"
    ws["B3"], ws["C3"], ws["D3"] = "Miesiąc", "Cel marży", "Aktualna marża"
    dane = [("Wrzesień 2026", 280000, 291000.25),
            ("Październik 2026", 1, 1),                 # duplikat – wygrywa wiersz niżej
            ("Październik 2026", 300000, 120000.5),
            ("Listopad 2026", "=C6*1.1", None),         # wiersz 7: formuła, Excel zapisał wynik 330000
            ("Grudzień 2026", "=C7*1.1", None),         # formuła bez zapisanego wyniku
            ("Razem", "=SUM(C4:C8)", "=SUM(D4:D8)"),
            ("styczeń 2027", 250000, 0)]
    for i, (lab, cel, akt) in enumerate(dane, start=4):
        ws.cell(row=i, column=2, value=lab)
        ws.cell(row=i, column=3, value=cel)
        ws.cell(row=i, column=4, value=akt)
    sw = wb.create_sheet("Święta PL")
    sw["A1"] = "Data"
    for i, d in enumerate(["2026-01-01", "2026-01-06", "2026-11-01", "2026-11-11", "2026-12-25", "2026-12-26",
                           "2027-01-01", "2027-01-06"], start=2):
        sw.cell(row=i, column=1, value=datetime.datetime.fromisoformat(d))
    wb.save(path)
    # wynik formuły w C7 zapisany tak, jak robi to Excel (<f>…</f><v>330000</v>) – openpyxl tego nie umie
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                xml2 = re.sub(r'(<c r="C7"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v>[^<]*</v>)?', r"\1\2<v>330000</v>", xml)
                assert xml2 != xml, "nie udało się dopisać wyniku formuły"
                data = xml2.encode("utf-8")
            out.writestr(item, data)
    os.replace(tmp, path)


def _uruchom(fail=False, prawdziwy_kalkulator=False, sesja=None):
    from marza import runner, podglad_wroclaw
    folder = tempfile.mkdtemp()
    if prawdziwy_kalkulator:
        _kalkulator(os.path.join(folder, "k.xlsx"))
    else:
        open(os.path.join(folder, "k.xlsx"), "wb").write(b"x" * 10)   # nie da się go odczytać jako Excel
    open(os.path.join(folder, "r.pdf"), "wb").write(b"%PDF")
    s = sesja or _Sesja()
    runner.SESSION = s
    # podgląd: prawdziwa funkcja odczytu z podglądu Wrocławia (bez podmiany atrybutów samego modułu)
    pod = types.SimpleNamespace(_read_source_months=podglad_wroclaw._read_source_months)
    mods = {"silnik_wroclaw": _silnik(folder, fail), "podglad_wroclaw": pod}
    old_load = runner._load
    runner._load = lambda name: mods[name]
    import smtplib as prawdziwy
    out = io.StringIO()
    try:
        with contextlib.redirect_stdout(out):
            rc = runner._engine(folder, "silnik_wroclaw", "podglad_wroclaw",
                                {"marza_recipients": {"wroclaw": ["a@example.com"]}}, lambda q: False)
    finally:
        runner._load = old_load
        runner.SESSION = None
    return rc, s.calls, sys.modules.get("smtplib") is prawdziwy, out.getvalue()


def test_przebieg_mail_przez_serwer_kopia_i_wyniki():
    os.environ.pop("LATEX_TRYB_TESTOWY", None)
    rc, calls, smtp_ok, log = _uruchom()
    assert rc == 0 and smtp_ok
    mail = [c for c in calls if c[1] == "/functions/v1/send-mail"]
    assert mail and mail[0][2]["action"] == "raw" and mail[0][2]["to"] == ["a@example.com"]
    assert any(c[0] == "UP" and c[1].endswith("/k.xlsx") for c in calls)
    assert any(c[0] == "UP" and c[1].endswith("/r.pdf") for c in calls)
    wynik = [c for c in calls if c[1] == "/rest/v1/marza_results"][0][2][0]
    assert wynik["margin"] == 12345.5 and wynik["target"] == 300000.0 and wynik["done"] == 120000.5
    # kalkulatora nie da się odczytać → tylko linia w dzienniku, przebieg OK, nic nie poszło do marza_months
    assert "Uwaga: wyniki dla zakładki WYNIKI nie zostały zapisane:" in log
    assert not [c for c in calls if c[1] == MIESIACE]


def test_awaria_silnika_przywraca_smtplib_i_nie_zapisuje_wynikow():
    os.environ.pop("LATEX_TRYB_TESTOWY", None)
    rc, calls, smtp_ok, _ = _uruchom(fail=True, prawdziwy_kalkulator=True)
    assert rc == 1 and smtp_ok
    assert not [c for c in calls if c[1] == "/rest/v1/marza_results"]
    assert not [c for c in calls if c[1] == MIESIACE]


def test_wyniki_miesiace_z_kalkulatora():
    os.environ.pop("LATEX_TRYB_TESTOWY", None)
    rc, calls, smtp_ok, log = _uruchom(prawdziwy_kalkulator=True)
    assert rc == 0 and smtp_ok
    zap = [c for c in calls if c[1] == MIESIACE]
    assert len(zap) == 1, zap
    m, _, rows, prefer = zap[0]
    assert m == "POST" and prefer == "resolution=merge-duplicates,return=minimal"
    by = {(r["year"], r["month"]): r for r in rows}
    assert sorted(by) == [(2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1)], sorted(by)   # bez „Razem”
    for r in rows:
        assert set(r) == {"branch", "year", "month", "label", "target", "done", "holidays", "updated_at"}
        assert r["branch"] == "wroclaw"
        assert r["target"] is None or type(r["target"]) is float
        assert r["done"] is None or type(r["done"]) is float
        datetime.datetime.fromisoformat(r["updated_at"])
    paz = by[(2026, 10)]
    assert paz["label"] == "Październik 2026" and paz["target"] == 300000.0 and paz["done"] == 120000.5
    assert paz["holidays"] == []
    lis = by[(2026, 11)]
    assert lis["target"] == 330000.0          # formuła → wynik zapisany przez Excel
    assert lis["done"] is None                # pusta komórka → None, nie 0
    assert lis["holidays"] == ["2026-11-01", "2026-11-11"]
    gru = by[(2026, 12)]
    assert gru["target"] is None              # formuła bez zapisanego wyniku → None, nie 0
    assert gru["holidays"] == ["2026-12-25", "2026-12-26"]
    sty = by[(2027, 1)]
    assert sty["label"] == "Styczeń 2027" and sty["target"] == 250000.0 and sty["done"] == 0.0
    assert sty["holidays"] == ["2027-01-01", "2027-01-06"]   # święta tylko z tego roku i miesiąca
    assert by[(2026, 9)]["done"] == 291000.25
    assert "Zakładka WYNIKI: zapisano 5 miesięcy (Wrocław)." in log
    assert "formuły bez wyliczonej wartości" in log and "(1)" in log


def test_wyniki_blad_sesji_nie_przerywa_przebiegu():
    os.environ.pop("LATEX_TRYB_TESTOWY", None)
    rc, calls, smtp_ok, log = _uruchom(prawdziwy_kalkulator=True, sesja=_Sesja(fail_paths=("/rest/v1/marza_months",)))
    assert rc == 0 and smtp_ok
    assert [c for c in calls if c[1] == "/rest/v1/marza_results"]
    assert "Uwaga: wyniki dla zakładki WYNIKI nie zostały zapisane: Brak połączenia z serwerem." in log
    assert "Zakładka WYNIKI: zapisano" not in log


def test_wyniki_zapisane_mimo_bledu_kopii_i_wynikow():
    os.environ.pop("LATEX_TRYB_TESTOWY", None)
    s = _Sesja(fail_paths=("/rest/v1/marza_results",), fail_upload=True)
    rc, calls, smtp_ok, log = _uruchom(prawdziwy_kalkulator=True, sesja=s)
    assert rc == 0 and smtp_ok
    assert "Uwaga: kopia w chmurze nie powiodła się" in log
    assert "Uwaga: zapis wyników w bazie nie powiódł się" in log
    assert [c for c in calls if c[1] == MIESIACE] and "Zakładka WYNIKI: zapisano 5 miesięcy (Wrocław)." in log


def test_tryb_testowy_nic_nie_wysyla_do_bazy():
    old = os.environ.get("LATEX_TRYB_TESTOWY")
    os.environ["LATEX_TRYB_TESTOWY"] = "1"
    try:
        rc, calls, smtp_ok, log = _uruchom(prawdziwy_kalkulator=True)
    finally:
        if old is None:
            os.environ.pop("LATEX_TRYB_TESTOWY", None)
        else:
            os.environ["LATEX_TRYB_TESTOWY"] = old
    assert rc == 0 and smtp_ok
    assert not [c for c in calls if c[0] == "UP" or c[1].startswith("/rest/v1/")], calls
    assert "[TRYB TESTOWY] Wyniki i pliki NIE zostały zapisane w chmurze." in log


def test_etykieta_miesiaca():
    from marza.runner import _label_ym, _liczba, _ile_miesiecy
    assert _label_ym("Październik 2026") == (2026, 10)
    assert _label_ym(" październik  2026 ") == (2026, 10)
    assert _label_ym("Styczeń 2027") == (2027, 1)
    for zla in ("Razem", "Październik", "Foo 2026", "Październik 26a", None, 2026, "Październik 1999",
                "Październik 2026²", "Październik 2026 r."):
        assert _label_ym(zla) is None, zla
    assert _liczba(300000) == 300000.0 and type(_liczba(300000)) is float
    assert _liczba("120 000,50") == 120000.5
    for brak in (None, "", "=SUM(C4:C8)", "abc", True, float("nan")):
        assert _liczba(brak) is None, brak
    assert [_ile_miesiecy(n) for n in (1, 2, 5, 12, 22)] == \
        ["1 miesiąc", "2 miesiące", "5 miesięcy", "12 miesięcy", "22 miesiące"]


def test_silniki_bez_adresow_w_kodzie():
    from pathlib import Path
    from _wspolne import APP
    for f in ("silnik_wroclaw.py", "silnik_opole.py"):
        s = (APP / "marza" / f).read_text(encoding="utf-8")
        assert "@" not in s.split("EMAIL_ODBIORCY = ")[1].split("\n")[0], f
        assert "latexopony" not in s and "latexserwis.pl" not in s, f
    assert Path(APP / "marza" / "runner.py").exists()
