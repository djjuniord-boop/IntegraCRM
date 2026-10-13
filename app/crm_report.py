"""Raport z CRM („Listy kontrolne” → eksport .xlsx/.csv): kolumny Data, Rejestracja, Oddział."""
import csv
from datetime import date, datetime

import plates

# Serwisy pokazywane zawsze w oknie (kolejne, nieznane, dochodzą z wczytanego pliku)
BRANCHES = ["Latex Wrocław", "Latex Opole", "Stangum Opole"]


def _key(s) -> str:
    return " ".join(str(s or "").strip().lower().split())


def _col(header, *names):
    h = [_key(x) for x in header]
    for n in names:
        if n in h:
            return h.index(n)
    return None


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _rows(path):
    low = str(path).lower()
    if low.endswith(".csv"):
        raw = open(path, "rb").read()
        try:
            txt = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = raw.decode("cp1250")
        dialect = csv.Sniffer().sniff(txt[:2000], delimiters=",;\t")
        return list(csv.reader(txt.splitlines(), dialect))
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Listy kontrolne"] if "Listy kontrolne" in wb.sheetnames else wb.worksheets[0]
    out = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def load(path) -> list[dict]:
    """Wiersze raportu: [{"date": date|None, "plate": str, "branch": str}]."""
    rows = _rows(path)
    hi = next((i for i, r in enumerate(rows[:10]) if _col(r, "rejestracja", "nr rejestracyjny") is not None), None)
    if hi is None:
        raise ValueError("W pliku nie ma kolumny „Rejestracja” – to nie jest raport z „Listy kontrolne” w CRM.")
    h = rows[hi]
    cp = _col(h, "rejestracja", "nr rejestracyjny")
    cd = _col(h, "data")
    cb = _col(h, "oddział", "oddzial", "serwis")
    out = []
    for r in rows[hi + 1:]:
        if cp >= len(r) or not r[cp]:
            continue
        plate = plates.normalize(str(r[cp]))
        if not plate:
            continue
        out.append({"date": _to_date(r[cd]) if cd is not None and cd < len(r) else None,
                    "plate": plate,
                    "branch": str(r[cb]).strip() if cb is not None and cb < len(r) and r[cb] else "(brak oddziału)"})
    if not out:
        raise ValueError("Raport z CRM nie zawiera żadnego numeru rejestracyjnego.")
    return out


def branches(rows) -> list[str]:
    seen = []
    for r in rows:
        if r["branch"] not in seen:
            seen.append(r["branch"])
    return seen


def select(rows, chosen) -> list[dict]:
    keys = {_key(b) for b in chosen}
    return [r for r in rows if _key(r["branch"]) in keys]
